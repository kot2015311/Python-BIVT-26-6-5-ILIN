"""ЛР5, задание B: конвертация CSV -> XLSX."""
import csv                                  # стандартный модуль для работы с CSV
from pathlib import Path                    # удобная работа с путями к файлам

from openpyxl import Workbook                       # Excel-книга (внешняя библиотека, см. requirements.txt)
from openpyxl.styles import Font                    # шрифт: сделаем заголовок жирным
from openpyxl.utils import get_column_letter        # номер колонки -> буква: 1 -> "A", 2 -> "B"

from src.lib.io_helpers import check_path           # проверка расширения и наличия файла
from src.lab04.io_txt_csv import ensure_parent_dir  # создание папок из ЛР4

MIN_WIDTH = 8   # минимальная ширина колонки в символах (по ТЗ)


def csv_to_xlsx(csv_path: str | Path, xlsx_path: str | Path) -> None:
    """
    Конвертирует CSV в XLSX через openpyxl.
    Первая строка CSV — заголовок (делаем жирным). Лист называется "Sheet1".
    Ширина колонки — по самому длинному тексту в ней, но не меньше 8 символов.
    """
    # Проверяем оба пути: вход должен быть .csv и существовать, выход — .xlsx
    src = check_path(csv_path, ".csv", must_exist=True)
    dst = check_path(xlsx_path, ".xlsx", must_exist=False)

    with src.open(encoding="utf-8", newline="") as f:
        # csv.reader отдаёт каждую строку списком; пустые строки ([]) пропускаем
        rows = [row for row in csv.reader(f) if row]

    # Пустой файл -> нет даже заголовка
    if len(rows) == 0:
        raise ValueError("Пустой CSV: нет заголовка")

    header = rows[0]
    # Пустое имя колонки — заголовок испорчен
    for name in header:
        if name.strip() == "":
            raise ValueError("В заголовке CSV есть пустое имя колонки")

    # Только заголовок без данных
    if len(rows) == 1:
        raise ValueError("Пустой CSV: есть только заголовок")

    # Все строки должны быть той же длины, что и заголовок, иначе таблица «поедет»
    for row in rows:
        if len(row) != len(header):
            raise ValueError("Все строки CSV должны иметь одинаковую длину.")

    wb = Workbook()             # новая книга Excel, в ней сразу есть один лист
    ws = wb.active              # берём этот активный лист
    ws.title = "Sheet1"         # имя листа по ТЗ

    # append добавляет список как новую строку листа: элементы -> ячейки A, B, C...
    for row in rows:
        ws.append(row)

    # ws[1] — первая строка листа (в Excel строки считаются с 1), делаем её жирной
    # Шрифт указываем явно: Font(bold=True) без name сбрасывает Calibri на шрифт по умолчанию
    for cell in ws[1]:
        cell.font = Font(name="Calibri", size=11, bold=True)

    # Автоширина: для каждой колонки ищем самый длинный текст
    for col in range(len(header)):
        longest = 0
        for row in rows:
            longest = max(longest, len(row[col]))
        # +2 — небольшой запас, чтобы текст не прилипал к границе; минимум 8 символов
        width = max(longest + 2, MIN_WIDTH)
        # column_dimensions["A"].width — ширина колонки A в символах
        ws.column_dimensions[get_column_letter(col + 1)].width = width

    # Создаём папку для файла и сохраняем книгу
    ensure_parent_dir(dst)
    wb.save(dst)
