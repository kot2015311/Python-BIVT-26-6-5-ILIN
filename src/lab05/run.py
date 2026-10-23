"""ЛР5: запуск трёх сценариев демонстрации и проверка результатов."""
import csv
import json

from openpyxl import load_workbook          # открыть готовый XLSX и посмотреть, что внутри

from src.lab05.json_csv import json_to_csv, csv_to_json
from src.lab05.csv_xlsx import csv_to_xlsx


# Сценарий 1: JSON -> CSV, проверяем заголовок и количество строк
json_to_csv("data/samples/people.json", "data/out/people_from_json.csv")
json_to_csv("data/samples/students.json", "data/out/students_from_json.csv")

with open("data/samples/people.json", encoding="utf-8") as f:
    objects = json.load(f)                  # сколько объектов было в JSON
with open("data/out/people_from_json.csv", encoding="utf-8", newline="") as f:
    reader = csv.DictReader(f)
    csv_rows = list(reader)                 # сколько строк данных получилось в CSV
    header = reader.fieldnames

print("1) JSON -> CSV")
print("   заголовок:", header)
print(f"   объектов в JSON: {len(objects)}, строк в CSV: {len(csv_rows)}")

# Сценарий 2: CSV -> JSON, сравниваем ключи и количество объектов
csv_to_json("data/samples/people.csv", "data/out/people_from_csv.json")

with open("data/samples/people.csv", encoding="utf-8", newline="") as f:
    reader = csv.DictReader(f)
    src_rows = list(reader)
    src_header = reader.fieldnames
with open("data/out/people_from_csv.json", encoding="utf-8") as f:
    result = json.load(f)

print("2) CSV -> JSON")
print("   колонки CSV:", src_header)
print("   ключи JSON: ", list(result[0].keys()))
print(f"   строк в CSV: {len(src_rows)}, объектов в JSON: {len(result)}")

# Сценарий 3: CSV -> XLSX, проверяем имя листа и ширину колонок
csv_to_xlsx("data/samples/people.csv", "data/out/people.xlsx")
csv_to_xlsx("data/samples/cities.csv", "data/out/cities.xlsx")

print("3) CSV -> XLSX")
for path in ("data/out/people.xlsx", "data/out/cities.xlsx"):
    wb = load_workbook(path)
    ws = wb.active
    # column_dimensions[буква].width — ширина, которую поставила csv_to_xlsx
    widths = {letter: dim.width for letter, dim in ws.column_dimensions.items()}
    print(f"   {path}: лист {ws.title!r}, строк {ws.max_row}, ширины {widths}")
