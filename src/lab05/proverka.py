"""ЛР5: проверка функций через assert (правильная работа и все ошибки из ТЗ)."""
import csv
import json

from openpyxl import load_workbook

from src.lab05.json_csv import json_to_csv, csv_to_json
from src.lab05.csv_xlsx import csv_to_xlsx

OUT = "data/out/proverka"   # сюда пишем временные результаты проверок
ERR = "data/samples/errors" # файлы с ошибками для проверок


def raises(error_type, func, *args) -> bool:
    """Возвращает True, если func(*args) упала именно с ошибкой error_type."""
    try:
        func(*args)
    except error_type:
        return True
    return False


# ---------- JSON -> CSV ----------
json_to_csv("data/samples/people.json", f"{OUT}/people.csv")
with open(f"{OUT}/people.csv", encoding="utf-8", newline="") as f:
    reader = csv.DictReader(f)
    rows = list(reader)
assert reader.fieldnames == ["name", "age", "city"], "Ошибка в json_to_csv (порядок колонок)"
assert len(rows) == 5, "Ошибка в json_to_csv (количество строк)"
assert rows[0]["name"] == "Alice" and rows[0]["age"] == "22", "Ошибка в json_to_csv (значения)"

json_to_csv("data/samples/students.json", f"{OUT}/students.csv")
with open(f"{OUT}/students.csv", encoding="utf-8", newline="") as f:
    reader = csv.DictReader(f)
    rows = list(reader)
assert reader.fieldnames == ["name", "group", "skills", "email"], "Ошибка в json_to_csv (объединение ключей)"
assert rows[0]["email"] == "", "Ошибка в json_to_csv (пустое поле вместо отсутствующего)"
assert rows[0]["skills"] == '["Python", "Git"]', "Ошибка в json_to_csv (вложенный список)"

# ---------- CSV -> JSON ----------
csv_to_json("data/samples/people.csv", f"{OUT}/people.json")
with open(f"{OUT}/people.json", encoding="utf-8") as f:
    data = json.load(f)
assert len(data) == 5, "Ошибка в csv_to_json (количество объектов)"
assert list(data[0].keys()) == ["name", "age", "city"], "Ошибка в csv_to_json (ключи)"
assert data[0]["age"] == "22", "Ошибка в csv_to_json (значения должны быть строками)"

# Туда и обратно: JSON -> CSV -> JSON даёт те же записи (возраст стал строкой)
csv_to_json(f"{OUT}/people.csv", f"{OUT}/round.json")
with open("data/samples/people.json", encoding="utf-8") as f:
    original = json.load(f)
with open(f"{OUT}/round.json", encoding="utf-8") as f:
    back = json.load(f)
assert back == [{k: str(v) for k, v in item.items()} for item in original], "Ошибка в обратимом преобразовании"

# ---------- CSV -> XLSX ----------
csv_to_xlsx("data/samples/cities.csv", f"{OUT}/cities.xlsx")
ws = load_workbook(f"{OUT}/cities.xlsx").active
assert ws.title == "Sheet1", "Ошибка в csv_to_xlsx (имя листа)"
assert ws.max_row == 6, "Ошибка в csv_to_xlsx (количество строк: заголовок + 5)"
assert ws["A1"].value == "city" and ws["A1"].font.bold, "Ошибка в csv_to_xlsx (заголовок)"
assert ws.column_dimensions["C"].width == 12, "Ошибка в csv_to_xlsx (ширина: len('population') + 2 = 12)"
assert ws.column_dimensions["A"].width >= 8, "Ошибка в csv_to_xlsx (ширина не меньше 8)"

# ---------- Ошибки ----------
# Нет файла -> FileNotFoundError
assert raises(FileNotFoundError, json_to_csv, "data/samples/nope.json", f"{OUT}/x.csv"), "Нет FileNotFoundError (json_to_csv)"
assert raises(FileNotFoundError, csv_to_json, "data/samples/nope.csv", f"{OUT}/x.json"), "Нет FileNotFoundError (csv_to_json)"
assert raises(FileNotFoundError, csv_to_xlsx, "data/samples/nope.csv", f"{OUT}/x.xlsx"), "Нет FileNotFoundError (csv_to_xlsx)"

# Неверный тип файла -> ValueError
assert raises(ValueError, csv_to_json, f"{ERR}/people.txt", f"{OUT}/x.json"), "Нет ValueError (вход .txt)"
assert raises(ValueError, json_to_csv, "data/samples/people.json", f"{OUT}/x.txt"), "Нет ValueError (выход .txt)"
assert raises(ValueError, csv_to_xlsx, "data/samples/people.csv", f"{OUT}/x.xls"), "Нет ValueError (выход .xls)"

# Пустой / неправильный JSON -> ValueError
assert raises(ValueError, json_to_csv, f"{ERR}/empty.json", f"{OUT}/x.csv"), "Нет ValueError (пустой JSON)"
assert raises(ValueError, json_to_csv, f"{ERR}/not_dicts.json", f"{OUT}/x.csv"), "Нет ValueError (не словари)"
assert raises(ValueError, json_to_csv, f"{ERR}/broken.json", f"{OUT}/x.csv"), "Нет ValueError (битый JSON)"

# Пустой / неправильный CSV -> ValueError
for func, out in ((csv_to_json, f"{OUT}/x.json"), (csv_to_xlsx, f"{OUT}/x.xlsx")):
    assert raises(ValueError, func, f"{ERR}/empty.csv", out), "Нет ValueError (пустой CSV)"
    assert raises(ValueError, func, f"{ERR}/header_only.csv", out), "Нет ValueError (только заголовок)"
    assert raises(ValueError, func, f"{ERR}/ragged.csv", out), "Нет ValueError (строка длиннее заголовка)"

print("Все проверки успешно пройдены! Функции работают правильно.")
